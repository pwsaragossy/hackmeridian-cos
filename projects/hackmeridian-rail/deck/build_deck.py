"""Build deck/radox-custody-rail.pptx from the organizers' DECK-TEMPLATE.pptx.

Run: python deck/build_deck.py   (needs python-pptx and Pillow), then
     soffice --headless --convert-to pdf --outdir deck deck/radox-custody-rail.pptx
Copy rules: PITCH.md Claims guard and [[wording-the-rio-miami-deck-readme-ef38]].
Every fact on a slide has its source in deck/REVIEW.md (claims table).
"""
import copy
import io
import math
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(ROOT, "DECK-TEMPLATE.pptx")
OUT = os.path.join(ROOT, "deck", "radox-custody-rail.pptx")
SHOT = lambda name: os.path.join(ROOT, "deck", "screens", name)

# Switch only on Paulo's written OK ([[zis-in-public-material-radox-is-bc05]]); then paste his approved wording.
BRIDGE = "Context: our explainer video on a Rio–Miami trade-corridor initiative; this project is its custody layer."

# LOOK.md (Signal adapted)
INK, PAPER, CARD, LINE_C = "0E0F11", "F7F6F2", "FFFFFF", "DFDDD6"
TEXT_D, SEC_D, LINE_D, PANEL_D = "E7E8EA", "A3A8B0", "2A2F36", "131517"
SEC_C, LIME, LIME_C, SLATE_D, SLATE_C = "4D5560", "C6F24E", "556F00", "9FB2CC", "3F5675"
TITLE, BODY, MONO = "Inter Tight", "Instrument Sans", "JetBrains Mono"

HASHES = [  # proofs/lifecycle.json, rio-miami-rail
    ("open", "operator", "2b3d51f0"),
    ("handover 1", "Warehouse → Port", "635c99ac"),
    ("handover 2", "Port → Carrier", "56a0f006"),
    ("handover 3", "Carrier → Customs", "ac268a49"),
    ("handover 4", "Customs → Recipient", "5199d4fb"),
    ("deliver", "Recipient → Holder", "919def8c"),
    ("burn", "existing redeem", "864fd164"),
]
CONTRACT = "CCBJVXMJYWPHTDJHDFPRGKPGXGEP24JOPJAMUS5C3BMZJTRTWR2H4LII"
N = 10

prs = Presentation(TEMPLATE)
COVER, CONTENT = prs.slides[0], prs.slides[1]
BLANK = prs.slide_layouts[0]


def rgb(h):
    return RGBColor.from_string(h)


def sid(shape):
    return shape.name.split(";")[1]


def clone(proto, keep, dark):
    """New slide with the template's fixed parts (bars, labels, footer, logo) copied from proto."""
    s = prs.slides.add_slide(BLANK)
    bg = copy.deepcopy(proto._element.cSld.bg)
    s._element.cSld.insert(0, bg)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = rgb(INK if dark else PAPER)
    for sh in proto.shapes:
        if sid(sh) not in keep:
            continue
        if sh.shape_type == 13:  # picture: re-add from its blob
            blob = sh.image.blob
            if dark and proto is CONTENT:  # dark logo on a dark slide: use the cover's light logo
                blob = next(x for x in COVER.shapes if sid(x) == "91").image.blob
            s.shapes.add_picture(io.BytesIO(blob), sh.left, sh.top, sh.width, sh.height)
        else:
            s.shapes._spTree.append(copy.deepcopy(sh._element))
    return s


def find(s, key):
    for sh in s.shapes:
        if ";" in sh.name and sid(sh) == key:
            return sh


def set_run_text(sh, text, color=None):
    r = sh.text_frame.paragraphs[0].runs[0]
    r.text = text
    if color:
        r.font.color.rgb = rgb(color)


def text(s, x, y, w, h, paras, size=26, color=INK, font=BODY, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.12, after=6):
    """paras: list of paragraphs; a paragraph is a str or a list of (text, overrides) runs."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = spacing
        para.space_after = Pt(after)
        for run_text, o in ([(p, {})] if isinstance(p, str) else p):
            r = para.add_run()
            r.text = run_text
            f = r.font
            f.name = o.get("font", font)
            f.size = Pt(o.get("size", size))
            f.bold = o.get("bold", bold)
            f.color.rgb = rgb(o.get("color", color))
    return tb


def box(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, weight=1.0, dash=False):
    b = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        b.adjustments[0] = min(0.12, 0.16 / min(w, h))
    if fill:
        b.fill.solid()
        b.fill.fore_color.rgb = rgb(fill)
    else:
        b.fill.background()
    if line:
        b.line.color.rgb = rgb(line)
        b.line.width = Pt(weight)
        if dash:
            b.line.dash_style = 4  # dash
    else:
        b.line.fill.background()
    b.shadow.inherit = False
    b._element.xpath(".//a:effectRef")[0].set("idx", "0")  # else LibreOffice draws the theme shadow
    return b


def dot(s, cx, cy, d, fill):
    return box(s, cx - d / 2, cy - d / 2, d, d, fill=fill, shape=MSO_SHAPE.OVAL)


def hline(s, x, y, w, color, t=0.025):
    return box(s, x, y - t / 2, w, t, fill=color)


def content_slide(n, pill, title, next_label, dark=False):
    keep = {"106", "107", "108", "109", "110", "121", "122", "123"}
    s = clone(CONTENT, keep, dark)
    set_run_text(find(s, "107"), f"PLATFORM {n:02d} / {N:02d}", LIME)
    set_run_text(find(s, "108"), "HACKMERIDIAN 2026 · LISBON")
    set_run_text(find(s, "109"), next_label, LIME if next_label != "FINAL STOP" else None)
    p = find(s, "110")
    set_run_text(p, pill, INK)
    p.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
    p.fill.fore_color.rgb = rgb(LIME)
    p.width = Inches(0.9 + 0.15 * len(pill))
    if dark:
        find(s, "106").fill.fore_color.rgb = rgb(PANEL_D)
        find(s, "121").fill.fore_color.rgb = rgb(LINE_D)
        set_run_text(find(s, "122"), find(s, "122").text_frame.text, SEC_D)
    text(s, 1.0, 2.2, 18.0, 1.1, [title], size=58, font=TITLE, bold=True, color=TEXT_D if dark else INK,
         spacing=1.0)
    return s


# ---------- 1. Cover: the rail, not the handover (owner, 2026-10-25) ----------
s = clone(COVER, {"84", "85", "86", "91", "92", "93", "94", "95", "96", "97", "98", "99", "100", "101"}, dark=True)
set_run_text(find(s, "85"), f"PLATFORM 01 / {N:02d}", LIME)
pill = find(s, "86")
set_run_text(pill, "TESTNET · TEST KEYS", LIME)
pill.line.color.rgb = rgb(LIME)
pill.width = Inches(0.9 + 0.15 * 19)
pill.text_frame.paragraphs[0].runs[0].font.size = Pt(18)
set_run_text(find(s, "100"), "#HackMeridian26", LIME)
for k in ("95", "96", "97"):
    find(s, k).fill.fore_color.rgb = rgb(LIME)
text(s, 1.0, 2.55, 12.5, 1.5, ["Radox Custody Rail"], size=100, font=TITLE, bold=True, color=LIME, spacing=0.95)
# Submission title and tagline, owner-approved 2026-10-26 ~01:30 UTC (README ec97106, FORM-ANSWERS.md).
text(s, 1.0, 4.15, 12.3, 0.9, ["SWIFT for goods in transit, on Stellar"], size=44, font=TITLE, bold=True, color=TEXT_D, spacing=1.0)
text(s, 1.0, 5.25, 12.3, 1.2, ["One shared custody record for goods in transit: every handover signed by both companies, checkable by anyone."],
     size=26, color=TEXT_D)
text(s, 1.0, 6.65, 12.5, 0.5, ["Pilot corridor: Rio–Miami"], size=24, color=SEC_D)
s.shapes.add_picture(SHOT("qr-page-dark.png"), Inches(15.75), Inches(1.75), Inches(2.6), Inches(2.6))
text(s, 14.3, 4.5, 5.5, 0.9, ["Scan: the live shipment", [("rio-miami-rail.vercel.app", {"font": MONO, "color": LIME, "size": 18})]],
     size=20, color=TEXT_D, align=PP_ALIGN.CENTER, after=2)
hline(s, 14.0, 6.15, 5.0, LIME_C, 0.03)
for i in range(6):
    dot(s, 14.0 + i * 1.0, 6.15, 0.26 if i in (0, 5) else 0.16, LIME)
text(s, 13.87, 6.42, 2.0, 0.4, ["RIO"], size=18, font=MONO, bold=True, color=SEC_D)
text(s, 17.13, 6.42, 2.0, 0.4, ["MIAMI"], size=18, font=MONO, bold=True, color=SEC_D, align=PP_ALIGN.RIGHT)

# ---------- 2-3. Problem, in numbers: only rows of PROBLEM-NUMBERS.md (verified 2026-10-25), wording kept ----------
# (big number, what it measures, source line). check_deck fails on any {{ placeholder.
P_HANDS = [("45M", "ocean bills of lading issued a year", "DCSA, Feb 2023"),  # PROBLEM-NUMBERS.md row 1, [[dcsa-45-million-bills-of-lading-1af2]]
           ("~11%", "of bills of lading issued electronically, 2025", "DCSA, as of 2025"),  # row 2
           ("22%", "of recorded cargo thefts in 2025 involved insiders", "TT Club & BSI Cargo Theft Report, Apr 2026")]  # row 3, [[tt-club-bsi-2025-cargo-theft-12f4]]
P_COST = [("R$900M", "est. direct loss, 8,570 cargo robberies in Brazil, 2025", "NTC&Logística survey, 28 Mar 2026"),  # row 7
          ("80–90%", "of world trade relies on trade finance", "WTO (undated page)"),  # Hin Leong cut: forgery by a signer is not stopped by dual signing (submission reviewer, 2026-10-26)
          ("US$2.5T", "global trade finance gap, unchanged from 2023", "ADB Global Trade Finance Gap Survey, Dec 2025")]  # [[adb-trade-finance-gap-survey-dec-d033]]


def stats(s, items, y, step=6.1, w=5.6):
    for i, (num, label, src) in enumerate(items):
        x = 1.0 + i * step
        hline(s, x, y, w, LIME_C, 0.03)
        text(s, x, y + 0.2, w, 1.2, [num], size=64, font=TITLE, bold=True, color=INK, spacing=0.95)
        text(s, x, y + 1.45, w, 1.0, [label], size=24, bold=True, color=INK)
        text(s, x, y + 2.45, w, 0.9, [src], size=18, color=SEC_C)


s = content_slide(2, "PMF · THE PROBLEM", "Who holds the whole chain?", "NEXT ▸ WHAT IT COSTS")
stops = ["Warehouse", "Port", "Carrier", "Customs", "Recipient"]
x0, x1, ly = 1.6, 18.4, 3.6
for i, name in enumerate(stops):  # broken line: each company is its own island
    cx = x0 + i * (x1 - x0) / 4
    if i < 4:
        hline(s, cx + 0.35, ly, (x1 - x0) / 4 - 0.7, LINE_C, 0.03)
    dot(s, cx, ly, 0.3, LIME_C)
    text(s, cx - 1.3, ly + 0.22, 2.6, 0.5, [name], size=22, bold=True, align=PP_ALIGN.CENTER)
text(s, 1.0, 4.45, 18.0, 0.5, [[("The gap: ", {"bold": True, "color": LIME_C}),
                               ("each company records only its own leg, and no shared record follows the goods.", {})]], size=22)
stats(s, P_HANDS, 5.15)
text(s, 1.0, 8.5, 18.0, 1.2, [
    [("Ports run legacy systems: ", {"bold": True, "color": INK}),
     ("over 90% of ports in low- and middle-income countries have no port community system (World Bank remark relayed by IAPH, Nov 2023); "
      "IMO guidance plans for legacy systems with proprietary formats (FAL.5/Circ.54, 2024).", {})],
    BRIDGE], size=18, color=SEC_C, after=4)

s = content_slide(3, "PMF · WHAT IT COSTS", "Cargo goes missing. Finance runs blind.", "NEXT ▸ THE STANDARD")
text(s, 1.0, 3.4, 18.0, 0.9, [[("The gap: ", {"bold": True, "color": LIME_C}),
                               ("a financier cannot see who holds the goods before paying.", {})]], size=24)
stats(s, P_COST, 4.6)
text(s, 1.0, 8.2, 18.0, 0.9, ["These figures size the market problem the rail addresses. Sources under each number."],
     size=20, color=SEC_C)

# ---------- 4. The proposal: a public SWIFT-like standard, integrate once; open vs operated ----------
s = content_slide(4, "THE PROPOSAL · A PUBLIC STANDARD", "A SWIFT-like standard: integrate once", "NEXT ▸ THE RAIL")
text(s, 1.0, 3.35, 18.0, 1.2, [[("The goal is ", {"bold": True}),
                                ("a shared, SWIFT-like custody standard that companies along a supply chain, from origin to destination, will use. ", {}),
                                ("We are building it on Stellar, starting with the custody rail.", {"bold": True})]], size=24, color=INK)


def link(x1, y1, x2, y2, color, w=1.5):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)


R = 0.95
for cx, hub, head in ((2.8, False, "PAIR BY PAIR · 10 LINKS"), (6.65, True, "ONE STANDARD · 5 LINKS")):
    cy = 5.95
    pts = [(cx + R * math.sin(2 * math.pi * k / 5), cy - R * math.cos(2 * math.pi * k / 5)) for k in range(5)]
    if hub:
        for x, y in pts:
            link(cx, cy, x, y, LIME_C, 2.0)
        dot(s, cx, cy, 0.42, LIME)
    else:
        for i in range(5):
            for j in range(i + 1, 5):
                link(*pts[i], *pts[j], SEC_C, 1.25)
    for x, y in pts:
        dot(s, x, y, 0.26, INK)
    text(s, cx - 1.8, 4.45, 3.6, 0.45, [head], size=18, font=MONO, bold=True, color=LIME_C if hub else SEC_C, align=PP_ALIGN.CENTER)
text(s, 1.0, 7.2, 7.8, 0.9, ["Example, 5 companies: 10 custom links, or 5 links to one standard."], size=20, color=SEC_C, align=PP_ALIGN.CENTER)
cols = [(9.3, "OPEN · ANYONE", LIME_C, ["Record format, roles and events: the custody_registry contract (Apache-2.0)", "Verify any shipment, no login (slide 6)"]),
        (14.35, "RADOX WILL OPERATE · CLIENT PAYS", INK, ["Onboarding, issuance controls, operation of the rail", "The API and backend: closed, not self-hostable", "Per-institution subscription + fee per shipment (slide 9)"])]
for x, head, hc, items in cols:
    text(s, x, 4.45, 4.8, 0.45, [head], size=18, font=MONO, bold=True, color=hc)
    hline(s, x, 4.98, 4.65, hc, 0.03)
    text(s, x, 5.15, 4.75, 3.0, [[("— ", {"color": hc}), (t_, {})] for t_ in items], size=21, after=8)
hline(s, 1.0, 8.3, 18.0, LINE_C, 0.02)
text(s, 1.0, 8.42, 18.0, 0.9, [
    "Each company integrates once instead of pair by pair. Radox is part of the Rio–Miami trade-corridor initiative and its technology rail: every port, airport and chamber that joins it will join the rail, through one adapter to the Radox API. Radox is in conversations with chambers of commerce in Brazil and Miami.",
    [("Radox is the rail, never the issuer. Regulation and licences stay with the client and its local counsel.", {"color": SEC_C})]],
    size=18, after=2)

# ---------- 5. The rail, end to end; the double signature is the rule of each link ----------
s = content_slide(5, "THE RAIL · LIFECYCLE BUILT", "One rail, origin to delivery", "NEXT ▸ LIVE ON TESTNET")
text(s, 1.0, 3.35, 18.0, 0.6, ["One record will follow the goods across company boundaries; port and airport systems will connect by API, not a rewrite."], size=24, color=INK)
nodes = [("Open", "operator", LIME_C), ("Warehouse", "", INK), ("Port", "", INK), ("Carrier", "", INK), ("Customs", "", INK),
         ("Recipient", "", INK), ("Delivered", "to the holder", LIME_C), ("Burned", "existing redeem¹", LIME_C)]
x0, x1, ly = 1.3, 18.7, 5.0
gap = (x1 - x0) / (len(nodes) - 1)
hline(s, x0, ly, x1 - x0, LIME_C, 0.04)
for i, (name, sub, c) in enumerate(nodes):
    cx = x0 + i * gap
    dot(s, cx, ly, 0.34, c)
    text(s, cx - 1.25, ly + 0.28, 2.5, 0.5, [name], size=21, bold=True, align=PP_ALIGN.CENTER)
    if sub:
        text(s, cx - 1.25, ly + 0.72, 2.5, 0.5, [sub], size=18, color=SEC_C, align=PP_ALIGN.CENTER)
    if 1 <= i <= 5:  # the four handovers plus delivery: each signed by both sides (proofs: "both signed")
        text(s, cx + 0.15, ly - 0.62, gap - 0.3, 0.45, ["2 signatures"], size=18, font=MONO, color=LIME_C, align=PP_ALIGN.CENTER)
text(s, 1.0, 6.45, 18.0, 1.7, [
    [("— ", {"color": LIME_C}), ("Rule of every link: ", {"bold": True}), ("both sides sign one Stellar transaction (delivery: recipient and holder), or the call fails.", {})],
    [("— ", {"color": LIME_C}), ("Each step extends the on-chain ", {}), ("head", {"font": MONO}), (": anyone can replay the history from Horizon and check it against the head.", {})],
    [("— ", {"color": LIME_C}), ("The token stays with its holder in transit; only custody moves.", {})]],
    size=23, after=6)
text(s, 1.0, 8.55, 18.0, 1.1, [
    [("Built: this lifecycle on testnet, test keys. Next: real companies on the rail, the standard, integrations.", {"bold": True, "color": INK})],
    "Keyed by sha256(container | origin); sits beside the carrier's bill of lading. ¹ Radox policy, not enforced on-chain; redeem needs the holder's signature."],
    size=18, color=SEC_C, after=2)

# ---------- 6. Proof: the whole chain ran on testnet ----------
s = content_slide(6, "THE PROOF · LIVE ON TESTNET", "The lifecycle ran on testnet", "NEXT ▸ WHY STELLAR", dark=True)
s.shapes.add_picture(SHOT("page-handovers-rows.png"), Inches(1.0), Inches(3.4), width=Inches(7.4))
box(s, 1.0, 3.4, 7.4, 7.4 * 950 / 1400, line=LINE_D)
tx = 9.0
for i, (step, who, h) in enumerate(HASHES):
    y = 3.3 + i * 0.37
    text(s, tx, y, 10.0, 0.4, [[(h, {"color": LIME, "font": MONO}), (f"   {step} · {who}", {"color": TEXT_D})]], size=18)
for i, (num, label) in enumerate([("7", "successful"), ("9", "refused, e.g. wrong signer (simulated)"), ("24", "unit tests"), ("MATCH", "head check, 25 Oct")]):
    x = tx + i * 2.55
    text(s, x, 5.95, 2.5, 0.8, [num], size=44, font=TITLE, bold=True, color=LIME, spacing=0.9)
    text(s, x, 6.72, 2.45, 0.7, [label], size=18, color=TEXT_D, spacing=1.0, after=0)
text(s, tx, 7.5, 10.0, 2.0, [
    [("custody_registry ", {"color": SEC_D}), (CONTRACT, {"color": TEXT_D})],
    [("Also on testnet: handovers signed from separate computers (proofs/three-machine.json); two test companies' chains on one registry, sharing stops, each dual-signed (proofs/multi-chain.json).", {"color": SEC_D, "font": BODY})],
    [("Testnet, test keys.", {"color": TEXT_D, "font": BODY, "bold": True})]],
    size=18, font=MONO, after=4)

# ---------- 7. Why Stellar ----------
s = content_slide(7, "ECOSYSTEM · WHY STELLAR", "Why Stellar: all signers, one call", "NEXT ▸ BUILT THIS WEEKEND")
rows = [("01", "Multi-party auth in one call", "Soroban authorises both companies of a link inside one contract invocation: the handover is atomic."),
        ("02", "History from Horizon, checked by an on-chain head", "RPC keeps contract events about 7 days, so we replay each shipment's transactions from Horizon and check them against the head stored in the contract."),
        ("03", "Any Stellar asset", "The record stores a token address and never calls the token, so any tokenized Stellar asset can carry a custody record.")]
for i, (n_, name, desc) in enumerate(rows):
    y = 3.55 + i * 1.5
    hline(s, 1.0, y - 0.12, 18.0, LINE_C, 0.02)
    text(s, 1.0, y + 0.08, 1.0, 0.6, [n_], size=24, font=MONO, bold=True, color=LIME_C)
    text(s, 2.0, y, 6.4, 1.2, [name], size=30, bold=True, font=TITLE)
    text(s, 8.6, y + 0.02, 10.4, 1.3, [desc], size=22, color=INK)
hline(s, 1.0, 7.93, 18.0, LINE_C, 0.02)
text(s, 1.0, 8.05, 18.0, 1.6, [
    [("Stellar holds about US$4B in tokenized assets ", {"bold": True}), ("(SDF, as of 1 Sep 2026). ", {"color": SEC_C}),
     ("Nearest neighbours ", {"bold": True}), ("(LumenLoop, SCF, Scout, 24–25 Oct): ", {"color": SEC_C}),
     ("Kutana and Trustless Work release escrow on trade milestones; Pipeline finances commodity trade. None records a handover signed by two custodians; no SEP.", {})],
    [("Open source, Apache-2.0: ", {"bold": True, "color": LIME_C}), ("a custody primitive any Stellar tokenized-asset issuer can reuse. Radox will open API access to them, end to end.", {})]],
    size=19, after=4, spacing=1.05)

# ---------- 8. Hackathon progress ----------
s = content_slide(8, "HACKATHON PROGRESS", "Built this weekend, on disclosed prior work", "NEXT ▸ MARKET AND ASK")
cols = [(1.0, "PRIOR WORK · BEFORE SUN 25 OCT 09:00", SEC_C, [
            "Radox platform (closed backend); its 3 Soroban contracts are on Stellar mainnet, public since 2026-08-19",
            "transfer_agent and its redeem (new testnet instance from the existing wasm)",
            "Architecture notes (20 Oct), planning files (24–25 Oct), the explainer video, our agent tooling (from 22 Oct)",
            "No contract code before then: first repo commit 25 Oct 09:00 UTC"]),
        (10.3, "BUILT AT HACKMERIDIAN", LIME_C, [
            "custody_registry contract, 24 unit tests",
            "Lifecycle harness: 7 testnet transactions, 9 refused calls",
            "verify-head script, the public shipment page, this deck",
            "Radox API custody route reads the registry, read-only: same head (testnet)"])]
for x, head, hc, items in cols:
    text(s, x, 3.5, 8.7, 0.5, [head], size=18, font=MONO, bold=True, color=hc)
    hline(s, x, 4.08, 8.7, hc, 0.03)
    text(s, x, 4.25, 8.7, 3.7, [[("— ", {"color": hc}), (t, {})] for t in items], size=21, after=8)
box(s, 1.0, 8.0, 18.0, 1.45, fill=CARD, line=LINE_C, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 1.35, 8.12, 17.3, 1.2, [[("AI attribution  ", {"font": MONO, "bold": True, "color": LIME_C, "size": 18}),
                                ("Code, tests, scripts, the shipment page, docs and this deck were made with AI coding agents (Claude Code); the team decided scope and what ships. Per-file record in AI_USAGE.md. The video narration is an AI voice generated from Pedro's voice.", {})]],
     size=20, anchor=MSO_ANCHOR.MIDDLE)

# ---------- 9. Market (MARKET-NUMBERS.md section 4, word for word), model (owner, 2026-10-25), next, ask ----------
# Rules: market figures, not ours; one trade-finance gap per slide (none here, slide 3 has ADB); no PortMiami TEU;
# pricing model with no values, Swift wording per [[swift-fee-structure-on-swift-com-9ed6]]; no money flow.
s = content_slide(9, "MARKET · MODEL · THE ASK", "One rail, every corridor", "NEXT ▸ THE TEAM")
text(s, 1.0, 3.3, 10.6, 0.4, ["MARKET · SOURCED FIGURES"], size=18, font=MONO, bold=True, color=LIME_C)
hline(s, 1.0, 3.78, 10.6, LIME_C, 0.03)
for i, (num, label) in enumerate([("US$94.2B", "US–Brazil goods trade, 2025¹"),
                                  ("US$6.1B", "Florida's goods exports to Brazil, its #1 market, 2025²")]):
    x = 1.0 + i * 5.4
    text(s, x, 3.95, 5.1, 0.85, [num], size=48, font=TITLE, bold=True, spacing=0.95)
    text(s, x, 4.8, 5.1, 0.8, [label], size=20, bold=True)
text(s, 1.0, 5.72, 10.6, 0.8, [[("Global: ", {"bold": True}),
     ("US$26.26T world merchandise exports, 2025³; ~80% of goods trade by volume goes by sea⁴; "
      "~45M ocean bills of lading a year⁵.", {})]], size=20)
text(s, 12.4, 3.3, 6.6, 0.4, ["MODEL"], size=18, font=MONO, bold=True, color=INK)
hline(s, 12.4, 3.78, 6.6, INK, 0.03)
text(s, 12.4, 3.95, 6.6, 1.4, ["Subscription per connected institution + a fee per shipment."], size=26, font=TITLE,
     bold=True, spacing=1.0)
text(s, 12.4, 5.25, 6.6, 1.4, ["Borrows the shape of Swift's billing rules: entry, annual, per-message usage⁶. "
                               "Not affiliated with Swift."], size=20, color=SEC_C)
text(s, 1.0, 6.6, 18.0, 0.45, [[("Now ", {"font": MONO, "bold": True, "color": LIME_C}),
     ("running on testnet, test keys.  ", {}),
     ("Next ", {"font": MONO, "bold": True, "color": LIME_C}),
     ("the Rio–Miami pilot shipment, SCF Build, audit, mainnet, API for other Stellar assets.", {})]], size=20)
text(s, 1.0, 7.12, 18.0, 0.7, ["¹ US Census Bureau, Trade in Goods with Brazil, 2025 (exports + imports, sum ours)  "
     "² USTR, Florida trade data, 2025  ³ WTO, Mar 2026  ⁴ UNCTAD, Review of Maritime Transport (undated page, "
     "opened 25 Oct 2026)  ⁵ DCSA, Feb 2023  ⁶ swift.com billing rules, accessed 25 Oct 2026"], size=18, color=SEC_C,
     spacing=1.0, after=0)
box(s, 1.0, 7.9, 18.0, 1.35, fill=LIME, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, 1.45, 7.9, 17.1, 1.35, [[("Our ask  ", {"font": MONO, "size": 20, "bold": True}),
                                  ("Introductions to one financial institution and one port, airport or logistics operator for a first pilot shipment, and a Stellar engineer's review of the design.", {"bold": True})]],
     size=24, anchor=MSO_ANCHOR.MIDDLE)

# ---------- 10. Team ----------
s = content_slide(10, "THE TEAM", "Why us: a founder already on Stellar mainnet", "FINAL STOP")
team = [("Pedro Wakigawa Saragossy", "Founder, captain. Built Radox: 3 Soroban contracts on Stellar mainnet."),
        ("Leo Schlanger", "Developer: custody_registry, the testnet harness and the proofs."),
        ("Vhendala", "QA and product review: reads every deliverable as a judge would.")]
for i, (name, role) in enumerate(team):
    y = 3.4 + i * 1.25
    text(s, 1.0, y, 10.8, 1.3, [[(name, {"bold": True, "size": 30, "font": TITLE})], [(role, {})]], size=21, color=INK, after=4)
text(s, 1.0, 7.15, 10.6, 0.5, ["Leo and Vhendala joined for the hackathon."], size=20, color=SEC_C)
if os.path.exists(SHOT("team.jpg")):  # drop the photo in deck/screens/team.jpg and rebuild
    s.shapes.add_picture(SHOT("team.jpg"), Inches(12.4), Inches(3.4), Inches(6.6), Inches(4.0))
else:
    box(s, 12.4, 3.4, 6.6, 4.0, line=SEC_C, shape=MSO_SHAPE.ROUNDED_RECTANGLE, dash=True)
    text(s, 12.4, 3.4, 6.6, 4.0, ["Team photo"], size=22, color=SEC_C, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 1.0, 7.65, 18.0, 0.6, ["Rio–Miami first. Our goal: one rail for goods, every corridor."], size=28, font=TITLE, bold=True, color=LIME_C)
hline(s, 1.0, 8.35, 18.0, LINE_C, 0.02)
text(s, 1.0, 8.45, 18.0, 1.2, [
    [("Live page  ", {"bold": True, "color": LIME_C}), ("rio-miami-rail.vercel.app", {}), ("     Code  ", {"bold": True, "color": LIME_C}),
     ("github.com/pwsaragossy/rio-miami-rail (Apache-2.0)", {})],
    [("AI  ", {"bold": True, "color": LIME_C}), ("Code, page, docs and deck made with AI coding agents (Claude Code). Narration: an AI voice from Pedro's voice. No AI-generated images in this deck.", {"font": BODY})]],
    size=18, font=MONO, after=6)

# drop the template's five original slides
ids = prs.slides._sldIdLst
for sld in list(ids)[:5]:
    prs.part.drop_rel(sld.rId)
    ids.remove(sld)
prs.save(OUT)
print("saved", OUT, len(prs.slides), "slides")
